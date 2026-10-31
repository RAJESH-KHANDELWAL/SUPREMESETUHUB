"""Normalize an official NIC 2025 workbook supplied by the operator.

Usage: python -m ...nic2025 INPUT.xlsx OUTPUT.json
Requires openpyxl. The source workbook is never modified.
"""
import hashlib
import json
import sys
from pathlib import Path

EXPECTED_LEVELS = ("section", "division", "group", "class", "subclass")

def normalize_workbook(source_path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl to import the official NIC 2025 workbook") from exc
    path = Path(source_path)
    raw = path.read_bytes()
    digest = hashlib.sha256(raw).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    records = []
    for sheet in workbook.worksheets:
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if not headers:
            continue
        normalized_headers = [str(value or "").strip().casefold() for value in headers]
        for row in rows:
            values = {normalized_headers[i]: value for i, value in enumerate(row) if i < len(normalized_headers) and normalized_headers[i]}
            if any(value is not None for value in row):
                records.append({"source_sheet": sheet.title, "source_values": values})
    workbook.close()
    if not records:
        raise ValueError("No non-empty rows found; inspect the official workbook sheet layout before mapping columns")
    return {"classification_id": "india-nic-2025", "source_file": path.name, "sha256": digest, "record_count": len(records), "records": records, "normalization_status": "raw_rows_extracted_column_mapping_review_required", "warning": "Column semantics and official hierarchy have not been inferred. Review headers and map codes/labels before publishing."}

def main():
    if len(sys.argv) != 3:
        raise SystemExit("Usage: python nic2025.py INPUT.xlsx OUTPUT.json")
    result = normalize_workbook(sys.argv[1])
    Path(sys.argv[2]).write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + "\\n", encoding="utf-8")
    print(f"Extracted {result['record_count']} rows; SHA-256 {result['sha256']}")

if __name__ == "__main__":
    main()
